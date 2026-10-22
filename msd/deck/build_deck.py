"""DAQ PCB troubleshooting deck (plain style), python-pptx."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from PIL import Image

INK, MUT, ACC, LINE = RGBColor(0x1a, 0x1a, 0x1a), RGBColor(0x55, 0x55, 0x55), RGBColor(0x1f, 0x4e, 0x9c), RGBColor(0xcc, 0xcc, 0xcc)
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
W, H = 13.333, 7.5
N = [0]


def para(tf, text, size=16, bold=False, color=INK, first=False, level=0, space=6):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.level = level
    p.space_after = Pt(space)
    parts = text.split("**")
    for i, part in enumerate(parts):
        if not part:
            continue
        r = p.add_run(); r.text = part
        r.font.size = Pt(size); r.font.bold = bold or (i % 2 == 1); r.font.color.rgb = color; r.font.name = "Calibri"
    return p


def slide(title, sub=None):
    s = prs.slides.add_slide(BLANK)
    N[0] += 1
    tb = s.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(W - 1.2), Inches(0.8))
    tf = tb.text_frame; tf.word_wrap = True
    para(tf, title, 28, True, first=True, space=0)
    if sub:
        para(tf, sub, 15, color=MUT, space=0)
    ln = s.shapes.add_connector(1, Inches(0.6), Inches(1.35 if sub else 1.15), Inches(W - 0.6), Inches(1.35 if sub else 1.15))
    ln.line.color.rgb = LINE; ln.line.width = Pt(1)
    nb = s.shapes.add_textbox(Inches(W - 1.0), Inches(H - 0.45), Inches(0.6), Inches(0.3))
    p = nb.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
    r = p.add_run(); r.text = str(N[0]); r.font.size = Pt(10); r.font.color.rgb = MUT
    return s


def bullets(s, items, x, y, w, h, size=16):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    for i, it in enumerate(items):
        lvl = 1 if it.startswith("  ") else 0
        para(tf, ("– " if lvl else "• ") + it.strip(), size - 2 * lvl, first=(i == 0), level=lvl)
    return tb


def pic(s, path, x, y, w=None, h=None):
    """Place an image inside the box (x, y, w, h), keeping its aspect ratio."""
    iw, ih = Image.open(path).size
    if w and h:
        if iw / ih > w / h:
            h2 = w * ih / iw; y += (h - h2) / 2; h = h2
        else:
            w2 = h * iw / ih; x += (w - w2) / 2; w = w2
    return s.shapes.add_picture(path, Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)


def table(s, rows, x, y, w, colw=None, size=13, rowh=0.38):
    nr, nc = len(rows), len(rows[0])
    t = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(rowh * nr)).table
    if colw:
        for j, cw in enumerate(colw):
            t.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.fill.solid(); c.fill.fore_color.rgb = RGBColor(0xee, 0xf1, 0xf6) if i == 0 else RGBColor(0xff, 0xff, 0xff)
            tf = c.text_frame; tf.word_wrap = True
            tf.paragraphs[0].text = ""
            para(tf, str(val), size, bold=(i == 0), first=True, space=0)
            c.margin_left = c.margin_right = Inches(0.08); c.margin_top = c.margin_bottom = Inches(0.04)
    return t


def note(s, text, y=6.85):
    tb = s.shapes.add_textbox(Inches(0.6), Inches(y), Inches(W - 2), Inches(0.4))
    tb.text_frame.word_wrap = True
    para(tb.text_frame, text, 11, color=MUT, first=True, space=0)


# 1 title -----------------------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK); N[0] += 1
tb = s.shapes.add_textbox(Inches(0.9), Inches(2.4), Inches(11.5), Inches(2.5)); tf = tb.text_frame; tf.word_wrap = True
para(tf, "DAQ PCB troubleshooting", 40, True, first=True)
para(tf, "MSD P25937 whisker board: ESP32 + ADS131M04 ADC + 2 × CD4097B mux + 32 INA351", 20, color=MUT)
para(tf, "Bench work 2026-10-19 to 2026-10-22", 16, color=MUT)

# 2 board -------------------------------------------------------------------------------------------
s = slide("The signal chain (one whisker, two gauges)",
          "gauge → quarter bridge → INA351 (G = 10, + REF) → RC → CD4097B mux → ADS131M04 → ESP32 → USB → host")
pic(s, "img_chain.png", 0.6, 1.55, 12.1, 5.2)
note(s, "8 mux selects × 4 ADC channels = 32 gauges per scan, ≈ 76 scans/s. Output ≈ REF + 23 mV per Ω of gauge above 352 Ω.")

# 3 overview -----------------------------------------------------------------------------------------
s = slide("Summary")
table(s, [["Symptom", "Cause", "Fix", "Status"],
          ["Spikes up to ~100 mV (W2_A)", "Firmware read stale conversions (DRDY level test)", "DRDY edge accounting + STATUS check", "fixed, 7.6 → 0.2 mV"],
          ["Rest level ≈ 0 V, bending clips", "REF ≈ 0 (see below); gauge ≈ 352 Ω sits at REF", "22 Ω inline resistors", "now counterproductive"],
          ["Offsets drift, ~1.1 V rest, hot-swap jumps", "REF floating: R7 bad joint, R8 cracked open", "R7 reflowed, R8 → 10 k through-hole", "REF 0.294 V, stable"],
          ["REF 0.29 V, not 0.59 V", "extra ~8–13 k load on REF", "–", "open, accepted"],
          ["Scan at 12.2 Hz", "ADC powered after ESP32 keeps OSR 1024 default", "STATUS RESET check + re-init", "firmware done"],
          ["CRC / stale storms", "ground wiring (loop, return path)", "one ground wire twisted with CLKIN", "fixed, CRC 0"],
          ["6–7 Hz tone, 0.2–0.3 mV", "845 Hz ripple on 3V3 × bridge imbalance, folded by the scan", "balance bridges (drop 22 Ω)", "explained"],
          ["W3_B, W8_B never rail", "mux 1 position 6 (X6 and Y6)", "–", "open"]],
      0.6, 1.4, 12.1, colw=[3.0, 4.0, 3.2, 1.9], size=13, rowh=0.55)

# 4 DRDY ------------------------------------------------------------------------------------------------
s = slide("Spiky noise → DRDY check in the firmware")
pic(s, "fig_drdy.png", 0.6, 1.3, 12.1, 4.2)
bullets(s, ["waitDRDY() tested whether DRDY was low instead of waiting for it to fall. A frame read (~240 µs) is shorter than a conversion (320 µs), so a scan could reuse an old conversion and the filter never flushed after a mux switch.",
            "Fix: count DRDY falling edges in an interrupt, require STATUS DRDY = 1111 on every frame (else count “stale”), keep the 4th fresh conversion after each mux switch."],
        0.6, 5.55, 12.1, 1.3, size=14)

# 5 22 ohm ------------------------------------------------------------------------------------------------
s = slide("Rest level ≈ 0 → inline 22 Ω resistors in the gauge leads")
bullets(s, ["OUT = 10 × (1.65 V − V_sense) + REF ≈ REF + 23 mV per Ω of gauge above 352 Ω (R3).",
            "A ~350 Ω gauge rests at REF. With REF ≈ 0 the INA351 cannot go below ~0.1 V, so one bending direction clipped.",
            "Workaround: 22 Ω in series with the gauge → rest ≈ +0.5 V above REF.",
            "Today REF is 0.29 V, so the headroom exists without them, and the extra imbalance costs noise (slides 13–14): the 22 Ω channels pick up ~10× more 3V3 ripple.",
            "Action: remove the 22 Ω where the gauge alone sits near 352 Ω."],
        0.6, 1.5, 12.1, 4.5, size=18)

# 6 floating REF ------------------------------------------------------------------------------------------
s = slide("Floating REF / offset voltage")
pic(s, "fig_floatref.png", 0.6, 1.3, 12.1, 4.0)
bullets(s, ["All 32 INA351 REF pins share one net; with no DC path REF floats on C4 (1 µF) and drifts.",
            "Symptoms: every channel drifting together, gauge + 22 Ω resting at 1.02–1.09 V instead of ~0.5 V, jumps after hot-swapping a whisker.",
            "Measured powered: REF = 3.28 V (pulled to 3V3) once R7 was reflowed and R8 turned out open."],
        0.6, 5.35, 12.1, 1.6, size=14)

# 7 two dividers ---------------------------------------------------------------------------------------------
s = slide("Two REF dividers in parallel; the fitted one had a broken R8")
pic(s, "img_map.png", 0.6, 1.35, 6.3, 4.2)
pic(s, "img_ref.png", 7.1, 1.35, 5.6, 4.2)
bullets(s, ["MAIN.SchDoc places the Amplifier_Reference sheet twice → two copies (top and bottom edge) on the one REF net. Positions add in parallel: either copy can set REF.",
            "Top copy fitted: R7 read 45 k or 100 k depending on probing (bad joint, reflowed → 46 k); R8 measured open off the board (cracked part). Bottom copy empty."],
        0.6, 5.65, 12.1, 1.4, size=14)

# 8 10k fix ---------------------------------------------------------------------------------------------------
s = slide("R8 replaced with a through-hole 10 k → REF stable at 0.29 V")
table(s, [["", "value"],
          ["R7 (3V3 → REF)", "≈ 46 k"],
          ["R8 (REF → GND)", "10 k (through-hole; spare from the same tape reads 10.0 k)"],
          ["REF expected", "3.3 V × 10 / (46 + 10) = 0.59 V"],
          ["REF measured", "0.294 V, stable"],
          ["R8 position, in circuit, unpowered", "5 k"]],
      0.6, 1.5, 7.4, colw=[3.0, 4.4], size=14, rowh=0.5)
bullets(s, ["Both readings say something else loads REF to GND:",
            "  from the voltage: 10 k ∥ Rx = 4.5 k → Rx ≈ 8 k",
            "  from the ohms: 10 k ∥ 46 k ∥ Rx = 5 k → Rx ≈ 13 k",
            "Not the INA351s (REF input 100 GΩ). No scorch marks; nearby via is GND. Unexplained, but stable, so accepted.",
            "Rest levels now as predicted: 370 Ω → 0.72 V (expected 0.70)."],
        8.3, 1.5, 4.5, 5.0, size=15)

# 9 12 Hz ---------------------------------------------------------------------------------------------------------
s = slide("Scan at 12.2 Hz: the ADC powered up after the ESP32")
pic(s, "fig_scan12.png", 0.6, 1.3, 12.1, 4.2)
bullets(s, ["The ESP32 configures the ADC (OSR 128) once at boot. Powered later (e.g. from the ROV 5 V), the ADC keeps its default OSR 1024: 2.56 ms per conversion × 4 per select × 8 selects = 81.92 ms per scan, exactly as logged.",
            "Fix in firmware: every frame's STATUS RESET bit is checked; a reset ADC is re-initialised automatically (counter adc_reinit). Side note: the 8× longer averaging was 2–5× quieter."],
        0.6, 5.55, 12.1, 1.4, size=14)

# 10 ground --------------------------------------------------------------------------------------------------------
s = slide("Ground wiring sets the CRC errors (and the noise)")
pic(s, "fig_ground.png", 0.6, 1.3, 7.4, 3.0)
table(s, [["Ground between ESP32 and board", "CRC errors", "scan"],
          ["two wires (MD1 pin 2 → both ESP32 GNDs), twisted with CLKIN / SCLK", "storm: 88 000 CRC, 14 500 stale", "66.7 Hz"],
          ["one loose wire (supply − → ESP32 GND)", "0 on 10-21; ~60 /s on 10-22, changed with wire position", "74 Hz"],
          ["one wire twisted with CLKIN", "0", "75.9 Hz"]],
      0.6, 4.45, 12.1, colw=[5.6, 4.8, 1.7], size=13, rowh=0.5)
bullets(s, ["Two ground paths form a loop through the ESP32 board.",
            "Twisting the only ground with CLKIN keeps the clock's return current next to it: small loop, little crosstalk onto SCLK.",
            "Rule: exactly one ground wire, run with the clock."],
        8.3, 1.3, 4.5, 3.0, size=14)

# 11 tone --------------------------------------------------------------------------------------------------------
s = slide("6–7 Hz noise in the scan data")
pic(s, "fig_tone.png", 0.6, 1.3, 12.1, 4.2)
bullets(s, ["~0.2–0.3 mV tone on the gauge and test-resistor channels; frequency wanders 6.1–7.0 Hz between runs.",
            "Unchanged with the laptop on battery and with another DC supply. Too regular to be the whisker: an alias of something faster."],
        0.6, 5.55, 12.1, 1.3, size=14)

# 12 burst --------------------------------------------------------------------------------------------------------
s = slide("Burst captures find the true frequency: 845 Hz")
bullets(s, ["New firmware command holds one mux select and records every conversion (8192 per burst). Three bursts at sample rates that are not multiples of each other leave only one frequency consistent with all three."],
        0.6, 1.4, 12.1, 1.0, size=15)
table(s, [["burst", "sample rate", "peak seen", "true frequency"],
          ["CLKIN 0.8 MHz, OSR 128", "3125 SPS", "845.05 Hz", "845 Hz"],
          ["CLKIN 1 MHz, OSR 128", "3906 SPS", "845.04 Hz", "845 Hz"],
          ["CLKIN 8 MHz, OSR 2048", "1953 SPS", "845.00 Hz", "845 Hz"],
          ["normal scan", "76.2 scans/s", "6.8 Hz", "845 − 11 × 76.2 = 6.8 Hz"]],
      0.6, 2.6, 9.0, colw=[3.0, 1.8, 1.8, 2.4], size=14, rowh=0.5)
bullets(s, ["A 0.5 % drift of the source (845 → 849 Hz) moves the folded tone from 6.8 to ~11 Hz: why it seemed to wander.",
            "Smaller lines on every channel (~0.02–0.05 mV): 125 Hz and harmonics, 437.5, 500, 1000 Hz."],
        0.6, 5.3, 12.1, 1.5, size=14)

# 12b burst spectra ----------------------------------------------------------------------------------
s = slide("Burst spectra: 845 Hz only where the bridge is off balance")
pic(s, "fig_burstspec.png", 0.6, 1.3, 12.1, 4.4)
bullets(s, ["Left: two test resistors recorded at the same instants. The 845 Hz line is 13× larger on the channel 13× further from balance.",
            "Right: two whisker gauges near balance. What is left is a set of small lines (125 Hz and harmonics, 437.5, 500, 1000 Hz) on a floor of ~0.9 µV/√Hz."],
        0.6, 5.8, 12.1, 1.2, size=14)

# 13 cable --------------------------------------------------------------------------------------------------------
s = slide("The 845 Hz follows the cable with the inline resistors")
pic(s, "fig_cable.png", 0.6, 1.3, 12.1, 4.0)
bullets(s, ["Both whiskers recorded at the same instants (select 2). Hands on the cable changed nothing: conducted, not picked up.",
            "Without the inline resistors (burst 3) the tone is gone on both channels."],
        0.6, 5.45, 12.1, 1.3, size=14)

# 14 imbalance ----------------------------------------------------------------------------------------------------
s = slide("Mechanism: 3V3 ripple × bridge imbalance")
pic(s, "fig_imbalance.png", 0.6, 1.3, 7.3, 5.3)
bullets(s, ["Both bridge arms (gauge/R3 and R1/R2) are fed from 3V3: ripple cancels in a balanced bridge and passes in proportion to the imbalance.",
            "tone ≈ (OUT − REF) × ripple / 3.3 V",
            "Two test resistors (burst 4) sit on one line through REF: 0.63–0.66 mV per V of offset → ~2.1 mV of 845 Hz on 3V3.",
            "The 22 Ω resistors put W6 0.47 V off balance → 0.31 mV. A balanced gauge: ~0.02 mV.",
            "Under bending the ripple scales with the signal (~0.07 % of the swing): harmless."],
        8.1, 1.4, 4.7, 5.4, size=14)

# 15 mux --------------------------------------------------------------------------------------------------------
s = slide("Open: two dead slots on mux 1 position 6")
pic(s, "fig_mux6.png", 0.6, 1.3, 7.3, 5.4)
bullets(s, ["An empty slot should rail at 1200 mV. W3_B (X6, pin 3) reads −110 to −120 mV, W8_B (Y6, pin 16) ~1–2 mV: their ADC input is not driven.",
            "X6 and Y6 share the chip's address decoder line 6, so one fault in the CD4097B explains both; two bad joints at once would be a coincidence.",
            "Check before reworking: hold select 6 (burst command) and measure pin 3 / pin 16 (INA side, ~3.2 V expected) and COM pins 1 / 17."],
        8.1, 1.4, 4.7, 5.4, size=14)

# 16 next ---------------------------------------------------------------------------------------------------------
s = slide("Open items and next steps")
bullets(s, ["Remove the 22 Ω inline resistors where the gauge alone sits near 352 Ω.",
            "Choose the scan settings for ~80 Hz with the lowest noise: CLKIN, OSR, SPI clock, settle, averaging per select. Evaluate candidates offline on the burst recordings first.",
            "Mux 1 position 6 (W3_B, W8_B): measure, then decide on rework or replacement.",
            "Unexplained ~8–13 k load on REF; source of the 845 Hz on 3V3 (needs a scope). Both low priority now.",
            "Keep the wiring fixed (one ground wire twisted with CLKIN) so noise measurements stay comparable."],
        0.6, 1.5, 12.1, 5.0, size=18)

prs.save("DAQ_PCB_troubleshooting.pptx")
print("saved", N[0], "slides")
